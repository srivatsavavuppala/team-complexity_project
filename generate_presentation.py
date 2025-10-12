from pptx import Presentation
from pptx.util import Pt


def add_title_slide(prs: Presentation, title_text: str, subtitle_text: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title_text
    subtitle = slide.placeholders[1]
    subtitle.text = subtitle_text


def add_bullets_slide(prs: Presentation, title_text: str, bullets: list[str]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title_text
    text_frame = slide.shapes.placeholders[1].text_frame
    text_frame.clear()
    for index, bullet in enumerate(bullets):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = bullet
        paragraph.level = 0
        paragraph.font.size = Pt(24)


def build_presentation() -> Presentation:
    prs = Presentation()

    # 1. Title
    add_title_slide(prs, "Project Presentation", "Subtitle / Tagline / Name")

    # 2. Introduction
    add_bullets_slide(
        prs,
        "Introduction",
        [
            "Purpose and context",
            "Audience and goals",
            "Scope and constraints",
        ],
    )

    # 3. Problem Overview
    add_bullets_slide(
        prs,
        "Problem Overview",
        [
            "Current state challenges",
            "Impact on stakeholders",
            "Why it matters now",
        ],
    )

    # 4. Key Problem 1
    add_bullets_slide(
        prs,
        "Key Problem 1",
        [
            "Symptom(s) observed",
            "Root cause hypothesis",
            "Evidence or metrics",
        ],
    )

    # 5. Key Problem 2
    add_bullets_slide(
        prs,
        "Key Problem 2",
        [
            "Symptom(s) observed",
            "Root cause hypothesis",
            "Evidence or metrics",
        ],
    )

    # 6. Solution Overview
    add_bullets_slide(
        prs,
        "Solution Overview",
        [
            "Proposed approach",
            "How it addresses problems",
            "Key benefits",
        ],
    )

    # 7. Solution Details
    add_bullets_slide(
        prs,
        "Solution Details",
        [
            "Architecture or workflow",
            "Technology choices",
            "Risks and mitigations",
        ],
    )

    # 8. Implementation Plan
    add_bullets_slide(
        prs,
        "Implementation Plan",
        [
            "Milestones and timeline",
            "Owners and resources",
            "Success criteria",
        ],
    )

    # 9. Conclusion
    add_bullets_slide(
        prs,
        "Conclusion",
        [
            "Recap of value",
            "Expected outcomes",
            "Call to action",
        ],
    )

    # 10. Thank You
    add_bullets_slide(
        prs,
        "Thank You",
        [
            "Questions?",
            "Contact information",
        ],
    )

    return prs


def main() -> None:
    prs = build_presentation()
    output_path = "/workspace/presentation_10_slides.pptx"
    prs.save(output_path)


if __name__ == "__main__":
    main()
